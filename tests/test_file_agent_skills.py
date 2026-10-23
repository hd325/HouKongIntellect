"""Real local document roundtrips and capability boundary checks."""
from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

import fitz
from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook
from pptx import Presentation

import tests  # Establish synthetic-only storage before application imports.
from services.file_agent.skills import FileSkillError, FileSkills, TOOL_DEFINITIONS


class FileSkillsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.skills = FileSkills({}, self.root / "out")

    def tearDown(self):
        self.temp.cleanup()

    def create(self, format, content):
        result = self.skills.execute("create_file", {"format": format, "filename": "result." + format, "content": content})
        self.assertTrue(result["validation"]["valid"])
        self.assertTrue(Path(result["path"]).is_file())
        return result

    def test_schema_and_unknown_capability(self):
        self.assertEqual({t["function"]["name"] for t in TOOL_DEFINITIONS},
                         {"inspect_file", "read_file", "create_file", "edit_file", "validate_file"})
        with self.assertRaisesRegex(FileSkillError, "无权"):
            self.skills.execute("read_file", {"file_id": "other-user-file"})
        with self.assertRaises(FileSkillError):
            self.skills.execute("run_python", {"code": "print(1)"})

    def test_unsupported_edit_options_fail_instead_of_being_silently_ignored(self):
        source = self.create("docx", {"paragraphs": ["Keep"]})
        with self.assertRaisesRegex(FileSkillError, "不能忽略"):
            self.skills.execute("edit_file", {"file_id": source["file_id"], "filename": "bad.docx",
                "operations": [{"op": "style_paragraph", "index": 0, "font_colour": "FF0000"}]})
        self.assertEqual(len(list((self.root / "out").glob("*.docx"))), 1)

    def test_font_changes_include_chinese_fonts_and_color_edit_retains_ppt_font(self):
        from docx.oxml.ns import qn
        word = self.create("docx", {"paragraphs": ["中文測試"]})
        edited = self.skills.execute("edit_file", {"file_id": word["file_id"], "filename": "font.docx",
            "operations": [{"op": "style_paragraph", "index": 0, "font_name": "SimSun"}]})
        run = Document(edited["path"]).paragraphs[0].runs[0]
        self.assertEqual(run._element.rPr.rFonts.get(qn("w:eastAsia")), "SimSun")
        ppt = self.create("pptx", {"slides": [{"title": "中文標題", "body": ["內容"]}]})
        first = self.skills.execute("edit_file", {"file_id": ppt["file_id"], "filename": "font.pptx",
            "operations": [{"op": "style_shape", "slide": 0, "shape": 0, "font_name": "SimSun"}]})
        second = self.skills.execute("edit_file", {"file_id": first["file_id"], "filename": "color.pptx",
            "operations": [{"op": "style_shape", "slide": 0, "shape": 0, "color": "FF0000"}]})
        run = Presentation(second["path"]).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
        self.assertEqual(run.font.name, "SimSun")
        self.assertEqual(run._r.rPr.find(qn("a:ea")).get("typeface"), "SimSun")
        excel = self.create("xlsx", {"sheets": [{"name": "Data", "rows": [["中文"]]}]})
        changed = self.skills.execute("edit_file", {"file_id": excel["file_id"], "filename": "font.xlsx",
            "operations": [{"op": "style_range", "sheet": "Data", "range": "A1", "font_name": "SimSun"}]})
        self.assertEqual(load_workbook(changed["path"]).active["A1"].font.name, "SimSun")

    def test_docx_targeted_edit_preserves_unrelated_runs_table_and_image(self):
        source = self.root / "source.docx"
        image = self.root / "pixel.png"
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
        pix.clear_with(0xFF)
        pix.save(image)
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Keep ").bold = True
        p.add_run("this formatting").italic = True
        p = doc.add_paragraph()
        p.add_run("Target ").bold = True
        p.add_run("paragraph").font.size = Pt(13)
        doc.add_table(rows=1, cols=2).cell(0, 0).text = "Unchanged cell"
        doc.add_picture(str(image))
        doc.save(source)
        original_bytes = source.read_bytes()
        self.skills.input_paths["src"] = source
        edited = self.skills.execute("edit_file", {"file_id": "src", "filename": "edited.docx",
            "operations": [{"op": "set_paragraph", "index": 1, "text": "Updated paragraph"}]})
        new = Document(edited["path"])
        self.assertEqual(new.paragraphs[1].text, "Updated paragraph")
        self.assertTrue(new.paragraphs[0].runs[0].bold)
        self.assertTrue(new.paragraphs[0].runs[1].italic)
        self.assertEqual(new.tables[0].cell(0, 0).text, "Unchanged cell")
        self.assertEqual(len(new.inline_shapes), 1)
        self.assertEqual(source.read_bytes(), original_bytes)
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(edited["path"]) as after:
            media = [n for n in before.namelist() if n.startswith("word/media/")]
            self.assertEqual({n: before.read(n) for n in media}, {n: after.read(n) for n in media})

    def test_calendar_pdf_read_to_xlsx_and_pdf(self):
        rows = [["2026-10-01", "National Day"], ["2026-12-25", "Christmas"]]
        calendar = self.create("pdf", {"title": "School calendar", "tables": [{"headers": ["Date", "Holiday"], "rows": rows}]})
        read = self.skills.execute("read_file", {"file_id": calendar["file_id"]})
        self.assertIn("2026-10-01", read["items"][0]["text"])
        self.assertFalse(read["has_more"])
        workbook = self.create("xlsx", {"sheets": [{"name": "Holidays", "headers": ["Date", "Holiday"], "rows": rows}]})
        wb = load_workbook(workbook["path"])
        try:
            self.assertEqual(wb.active["B3"].value, "Christmas")
        finally:
            wb.close()
        with fitz.open(calendar["path"]) as pdf:
            self.assertEqual(len(pdf), 1)
            self.assertIn("National Day", pdf[0].get_text().replace("\xa0", " "))

    def test_ten_page_ppt_and_continued_version(self):
        result = self.create("pptx", {"slides": [{"title": f"Review {i + 1}", "body": ["First teaching point", "Second teaching point"]} for i in range(10)]})
        first_read = self.skills.execute("read_file", {"file_id": result["file_id"], "limit": 3})
        self.assertTrue(first_read["has_more"])
        self.assertEqual(first_read["next_start"], 3)
        self.assertEqual(first_read["pagination"]["total"], 10)
        edited = self.skills.execute("edit_file", {"file_id": result["file_id"], "filename": "review-v2.pptx",
            "operations": [{"op": "replace_text", "old": "Review 1", "new": "Updated Review 1"}]})
        prs = Presentation(edited["path"])
        self.assertEqual(len(prs.slides), 10)
        self.assertEqual(prs.slides[0].shapes[0].text, "Updated Review 1")
        self.assertNotEqual(result["file_id"], edited["file_id"])
        self.assertTrue(Path(result["path"]).is_file())

    def test_xlsx_formula_style_and_untouched_sheet(self):
        result = self.create("xlsx", {"sheets": [
            {"name": "Scores", "headers": ["Score", "Total"], "rows": [[10, "=SUM(A2:A3)"], [20, None]]},
            {"name": "Keep", "rows": [["Original", "=1+1"]]}]})
        edited = self.skills.execute("edit_file", {"file_id": result["file_id"], "filename": "scores-v2.xlsx", "operations": [
            {"op": "set_cell", "sheet": "Scores", "cell": "B2", "value": "=AVERAGE(A2:A3)"},
            {"op": "style_range", "sheet": "Scores", "range": "A1:B1", "fill": "ABCDEF", "bold": True},
            {"op": "column_width", "sheet": "Scores", "column": "A", "width": 30}]})
        wb = load_workbook(edited["path"], data_only=False)
        try:
            self.assertEqual(wb["Scores"]["B2"].value, "=AVERAGE(A2:A3)")
            self.assertEqual(wb["Scores"]["A1"].fill.fgColor.rgb, "00ABCDEF")
            self.assertEqual(wb["Keep"]["B1"].value, "=1+1")
            self.assertEqual(wb["Scores"].column_dimensions["A"].width, 30)
        finally:
            wb.close()

    def test_safety_and_explicit_errors(self):
        for filename in ("../escape.docx", "C:\\escape.docx", "wrong.pdf"):
            with self.subTest(filename=filename), self.assertRaises(FileSkillError):
                self.skills.execute("create_file", {"format": "docx", "filename": filename, "content": {"paragraphs": ["hi"]}})
        with self.assertRaisesRegex(FileSkillError, "外部"):
            self.create("xlsx", {"sheets": [{"rows": [["=WEBSERVICE(\"https://example.com\")"]]}]})
        source = self.root / "broken.docx"
        source.write_bytes(b"not an office file")
        self.skills.input_paths["broken"] = source
        with self.assertRaisesRegex(FileSkillError, "损坏"):
            self.skills.execute("read_file", {"file_id": "broken"})
        result = self.create("docx", {"paragraphs": ["Original"]})
        with self.assertRaisesRegex(FileSkillError, "未找到"):
            self.skills.execute("edit_file", {"file_id": result["file_id"], "filename": "bad.docx", "operations": [{"op": "replace_text", "old": "missing", "new": "new"}]})
        self.assertFalse(any(p.name.endswith("_bad.docx") for p in self.skills.output_dir.iterdir()))

    def test_pdf_pagination_keeps_all_text_and_cjk(self):
        paragraphs = [f"第{i}项放假安排：学校假期内容，来源校历。" for i in range(120)]
        result = self.create("pdf", {"title": "学校放假日期", "paragraphs": paragraphs})
        self.assertGreater(result["validation"]["page_count"], 1)
        read = self.skills.execute("read_file", {"file_id": result["file_id"], "limit": 1})
        self.assertTrue(read["has_more"])
        all_text = ""
        start = 0
        while True:
            read = self.skills.execute("read_file", {"file_id": result["file_id"], "start": start, "limit": 1})
            all_text += "".join(item["text"] for item in read["items"])
            if not read["has_more"]:
                break
            start = read["next_start"]
        self.assertIn("第119项放假安排", all_text)

    def test_run_replacement_with_repeated_text_and_preserved_drawing(self):
        result = self.create("docx", {"paragraphs": ["cat cat", "Keep this"]})
        updated = self.skills.execute("edit_file", {"file_id": result["file_id"], "filename": "repeated.docx",
            "operations": [{"op": "replace_text", "old": "cat", "new": "catcat"}]})
        self.assertEqual(Document(updated["path"]).paragraphs[0].text, "catcat catcat")

    def test_ppt_only_target_slide_xml_changes(self):
        result = self.create("pptx", {"slides": [
            {"title": "First", "body": ["Keep"], "notes": "Private notes"},
            {"title": "Second", "body": ["Change"]}]})
        updated = self.skills.execute("edit_file", {"file_id": result["file_id"], "filename": "one-slide.pptx",
            "operations": [{"op": "set_shape_text", "slide": 1, "shape": 0, "text": "Updated"}]})
        with zipfile.ZipFile(result["path"]) as before, zipfile.ZipFile(updated["path"]) as after:
            self.assertEqual(before.read("ppt/slides/slide1.xml"), after.read("ppt/slides/slide1.xml"))
            self.assertEqual(before.read("ppt/notesSlides/notesSlide1.xml"), after.read("ppt/notesSlides/notesSlide1.xml"))
        read = self.skills.execute("read_file", {"file_id": updated["file_id"]})
        self.assertEqual(read["items"][0]["notes"], "Private notes")

    def test_pdf_edit_keeps_other_page_render_identical_and_refuses_overflow(self):
        result = self.create("pdf", {"pages": [
            {"title": "First page", "paragraphs": ["Original content"]},
            {"title": "Keep page", "paragraphs": ["Never change this"]}]})
        updated = self.skills.execute("edit_file", {"file_id": result["file_id"], "filename": "edited.pdf",
            "operations": [{"op": "replace_text", "page": 0, "old": "Original", "new": "Updated"}]})
        with fitz.open(result["path"]) as before, fitz.open(updated["path"]) as after:
            self.assertEqual(before[1].get_pixmap().samples, after[1].get_pixmap().samples)
            self.assertIn("Updated", after[0].get_text())
        with self.assertRaisesRegex(FileSkillError, "放不下"):
            self.skills.execute("edit_file", {"file_id": result["file_id"], "filename": "overflow.pdf",
                "operations": [{"op": "replace_text", "page": 0, "old": "Original", "new": "This will not fit " * 100}]})

    def test_markitdown_local_read_and_image_requires_ocr(self):
        source = self.root / "table.csv"
        source.write_text("Date,Holiday\n2026-10-01,National Day\n", encoding="utf-8")
        self.skills.input_paths["csv"] = source
        read = self.skills.execute("read_file", {"file_id": "csv"})
        self.assertEqual(read["markdown_converter"], "markitdown-local-no-plugins")
        self.assertIn("| Date | Holiday |", read["markdown"])
        image = self.root / "scan.png"
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 5, 20), False)
        pix.clear_with(255)
        pix.save(image)
        self.skills.input_paths["scan"] = image
        read = self.skills.execute("read_file", {"file_id": "scan"})
        self.assertTrue(read["items"][0]["needs_ocr"])
        self.assertEqual(read["items"][0]["text"], "")

    def test_pdf_active_content_is_rejected(self):
        source = self.root / "active.pdf"
        with fitz.open() as pdf:
            pdf.new_page()
            pdf.set_xml_metadata("<root />")
            action = pdf.get_new_xref()
            pdf.update_object(action, "<</S /JavaScript /JS (app.alert\\(1\\))>>")
            pdf.xref_set_key(pdf.pdf_catalog(), "OpenAction", f"{action} 0 R")
            pdf.save(source)
        self.skills.input_paths["active"] = source
        with self.assertRaisesRegex(FileSkillError, "脚本"):
            self.skills.execute("read_file", {"file_id": "active"})

    def test_images_are_real_and_replaceable_in_three_office_formats(self):
        def make_image(name, color):
            target = self.root / name
            pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 16, 16), False)
            pix.clear_with(color)
            pix.save(target)
            self.skills.input_paths[name] = target
            return target
        make_image("one.png", 255)
        second = make_image("two.png", 0)
        word = self.create("docx", {"paragraphs": ["Image"], "images": [{"file_id": "one.png"}]})
        word2 = self.skills.execute("edit_file", {"file_id": word["file_id"], "filename": "image-replaced.docx", "operations": [
            {"op": "replace_image", "index": 0, "file_id": "two.png"}, {"op": "add_equation", "text": "x² + y² = r²"}]})
        doc = Document(word2["path"])
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertIn("x²", doc.element.xml)
        rel_id = doc.inline_shapes[0]._inline.graphic.graphicData.pic.blipFill.blip.embed
        self.assertEqual(doc.part.rels[rel_id].target_part.blob, second.read_bytes())
        ppt = self.create("pptx", {"slides": [{"title": "Image", "body": ["Keep text"], "image": {"file_id": "one.png"}}]})
        ppt2 = self.skills.execute("edit_file", {"file_id": ppt["file_id"], "filename": "image-replaced.pptx", "operations": [
            {"op": "replace_image", "slide": 0, "shape": 2, "file_id": "two.png"}]})
        self.assertEqual(Presentation(ppt2["path"]).slides[0].shapes[2].image.blob, second.read_bytes())
        excel = self.create("xlsx", {"sheets": [{"name": "Images", "rows": [["Keep"]]}]})
        excel2 = self.skills.execute("edit_file", {"file_id": excel["file_id"], "filename": "image-added.xlsx", "operations": [
            {"op": "add_image", "sheet": "Images", "file_id": "one.png", "cell": "B2", "width_px": 32, "height_px": 32}]})
        excel3 = self.skills.execute("edit_file", {"file_id": excel2["file_id"], "filename": "image-replaced.xlsx", "operations": [
            {"op": "replace_image", "sheet": "Images", "index": 0, "file_id": "two.png"}]})
        wb = load_workbook(excel3["path"])
        try:
            self.assertEqual(len(wb.active._images), 1)
            self.assertEqual(wb.active._images[0]._data(), second.read_bytes())
            self.assertEqual(wb.active["A1"].value, "Keep")
        finally:
            wb.close()

    def test_source_formulas_and_word_external_fields_rejected_before_read(self):
        source = self.root / "unsafe.xlsx"
        from openpyxl import Workbook
        wb = Workbook()
        wb.active["A1"] = '=WEBSERVICE("https://example.invalid")'
        wb.save(source)
        wb.close()
        self.skills.input_paths["unsafe-excel"] = source
        with self.assertRaisesRegex(FileSkillError, "外部"):
            self.skills.execute("read_file", {"file_id": "unsafe-excel"})
        source = self.root / "unsafe.docx"
        doc = Document()
        from docx.oxml import OxmlElement
        field = OxmlElement("w:instrText")
        field.text = ' INCLUDETEXT "https://example.invalid" '
        doc.add_paragraph().add_run()._r.append(field)
        doc.save(source)
        self.skills.input_paths["unsafe-word"] = source
        with self.assertRaisesRegex(FileSkillError, "外部"):
            self.skills.execute("read_file", {"file_id": "unsafe-word"})

    def test_slide_text_overflow_is_error(self):
        with self.assertRaisesRegex(FileSkillError, "超出"):
            self.create("pptx", {"slides": [{"title": "Too much", "body": ["Very long teaching content. " * 10] * 15}]})


if __name__ == "__main__":
    unittest.main()
