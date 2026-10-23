"""Document-agent API contracts; temporary storage, synthetic accounts only."""
from __future__ import annotations

import asyncio
import copy
import io
import json
import shutil
import tempfile
import threading
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import AsyncMock, patch

from fastapi.testclient import TestClient
from docx import Document

import tests  # Establish synthetic-only storage before application imports.
import main
from tests.account_fixtures import fixture_accounts
import httpx
from services.file_agent.agent import DeepSeekFileAgent


class FileAgentApiTests(unittest.TestCase):
    def setUp(self):
        _accounts = patch.dict(main.VALID_ACCOUNTS, fixture_accounts())
        _accounts.start()
        self.addCleanup(_accounts.stop)
        self.temp = tempfile.TemporaryDirectory(prefix="hk_file_api_")
        self.root = Path(self.temp.name)
        self.uploads = self.root / "uploads"
        self.uploads.mkdir()
        self.patches = [patch.object(main, "DATA_DIR", str(self.root)),
                        patch.object(main, "UPLOAD_DIR", str(self.uploads)),
                        patch.object(main, "CONVERSATIONS_FILE", str(self.root / "conversations.json")),
                        patch.object(main, "SESSIONS_FILE", str(self.root / "sessions.json")),
                        patch("main.logging.info"), patch("main.logging.warning"), patch("main.logging.error")]
        for item in self.patches:
            item.start()
        self.store = main.DataStore()
        now = datetime.now().isoformat()
        self.store._sessions = {f"session-{user}": {"account": user, "created_at": now,
             "current_conversation_id": f"conversation-{user}"} for user in ("alice", "bob")}
        self.store._conversations = {f"conversation-{user}": {"id": f"conversation-{user}",
            "user_id": user, "created_at": now, "updated_at": now, "title": "test", "messages": []}
            for user in ("alice", "bob")}
        self.store._files = []
        self.store._deleted_document_ids = set()
        self.store._file_contents = {}
        self.store_patch = patch.object(main, "data_store", self.store)
        self.store_patch.start()
        self.client = TestClient(main.app)

    def tearDown(self):
        self.client.close()
        self.store_patch.stop()
        for item in reversed(self.patches):
            item.stop()
        self.temp.cleanup()

    def ask(self, question, ids=(), user="alice", **extra):
        return self.client.post("/api/chat", headers={"X-Session-ID": "session-" + user},
            json={"question": question, "session_id": "untrusted-body-session", "file_ids": list(ids), **extra})

    def add(self, file_id, path, owner="alice", text=""):
        self.store._files.append({"id": file_id, "name": path.name, "original_name": path.name,
            "storage_path": path.name, "owner": owner, "audience": "personal", "size": path.stat().st_size,
            "uploaded_at": datetime.now().isoformat(), "review_status": "published"})
        self.store._file_contents[file_id] = text

    def events(self, response):
        self.assertEqual(response.status_code, 200, response.text)
        return [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: ")]

    def model(self, steps):
        remaining = iter(steps)

        async def complete(agent, messages, tools):
            step = next(remaining, None)
            if step is None:
                return {"content": json.dumps({"status": "complete", "summary": "已完成。",
                    "output_file_ids": list(agent.runtime.generated), "unmet_requirements": []})}
            if callable(step):
                step = step(messages)
            name, arguments = step
            return {"content": "", "tool_calls": [{"id": "call-" + str(len(messages)), "type": "function",
                      "function": {"name": name, "arguments": json.dumps(arguments, ensure_ascii=False)}}]}
        return patch.object(DeepSeekFileAgent, "complete", complete)

    def test_no_authentication_never_runs_tools(self):
        response = self.client.post("/api/chat", json={"question": "生成 PDF", "session_id": "session-alice"})
        self.assertEqual(response.status_code, 401)

    def test_other_users_reference_is_rejected_before_execution(self):
        path = self.uploads / "secret.txt"
        path.write_text("BOB_PRIVATE_TEXT", encoding="utf-8")
        self.add("bob-secret", path, "bob")
        with patch.object(DeepSeekFileAgent, "complete", side_effect=AssertionError("must not call AI")):
            response = self.ask("把引用文件生成 PDF", ["bob-secret"])
        self.assertEqual(response.status_code, 403)
        self.assertNotIn("BOB_PRIVATE_TEXT", response.text)
        self.assertFalse((self.uploads / "file_agent").exists())

    def test_untrusted_body_session_does_not_select_identity(self):
        path = self.uploads / "secret.txt"
        path.write_text("secret", encoding="utf-8")
        self.add("secret", path, "bob")
        response = self.ask("修改 Word", ["secret"], session_id="session-bob")
        self.assertEqual(response.status_code, 403)

    def test_other_users_conversation_is_rejected(self):
        self.store._sessions["session-alice"]["current_conversation_id"] = "conversation-bob"
        response = self.ask("生成 Excel")
        self.assertEqual(response.status_code, 403)

    def test_missing_original_returns_explicit_error(self):
        self.store._files = [{"id": "missing", "name": "missing.docx", "owner": "alice", "audience": "personal"}]
        response = self.ask("修改 Word 指定段落", ["missing"])
        events = self.events(response)
        self.assertEqual(events[-1]["type"], "error")
        self.assertIn("原始", events[-1]["detail"])
        self.assertFalse(any(event["type"] == "artifacts" for event in events))

    def test_atomic_commit_rolls_back_all_files_and_message_on_disk_error(self):
        outputs = [({"id": "a", "owner": "alice", "audience": "personal"}, "content-a"),
                   ({"id": "b", "owner": "alice", "audience": "personal"}, "content-b")]
        original = copy.deepcopy(self.store._conversations["conversation-alice"])
        with patch.object(self.store, "_write_json", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                asyncio.run(self.store.commit_file_task("conversation-alice", "alice", outputs, {"content": "done"}))
        self.assertEqual(self.store._files, [])
        self.assertEqual(self.store._file_contents, {})
        self.assertEqual(self.store._conversations["conversation-alice"], original)

    def test_cancellation_during_disk_commit_cannot_diverge_memory_and_disk(self):
        writing, finish = threading.Event(), threading.Event()
        original_write = self.store._write_json
        def delayed(data):
            writing.set()
            finish.wait(5)
            original_write(data)
        async def exercise():
            outputs = [({"id": "cancel-output", "owner": "alice", "audience": "personal"}, "content")]
            task = asyncio.create_task(self.store.commit_file_task("conversation-alice", "alice", outputs,
                {"content": "done", "artifacts": [{"id": "cancel-output"}]}))
            await asyncio.to_thread(writing.wait, 5)
            task.cancel()
            finish.set()
            await task
        with patch.object(self.store, "_write_json", delayed):
            asyncio.run(exercise())
        disk = json.loads((self.root / "conversations.json").read_text(encoding="utf-8"))
        self.assertEqual([item["id"] for item in disk["files"]], ["cancel-output"])
        self.assertEqual([item["id"] for item in self.store._files], ["cancel-output"])

    def test_word_edit_download_preview_and_followup_new_version(self):
        from docx import Document
        source = self.uploads / "lesson.docx"
        document = Document()
        document.add_paragraph("KEEP: original first paragraph.")
        document.add_paragraph("Change this paragraph.")
        document.save(source)
        self.add("word-source", source)
        steps = [("read_file", {"file_id": "word-source"}),
                 ("edit_file", {"file_id": "word-source", "filename": "lesson-v2.docx",
                                "operations": [{"op": "set_paragraph", "index": 1, "text": "Updated paragraph."}]})]
        with self.model(steps):
            events = self.events(self.ask("修改 Word 第2段为 Updated paragraph.，其他不变", ["word-source"]))
        self.assertEqual(events[-1]["type"], "done", events)
        artifact = events[-1]["artifacts"][0]
        download = self.client.get(f"/api/files/{artifact['id']}/download", headers={"X-Session-ID": "session-alice"})
        self.assertEqual(download.status_code, 200)
        updated = Document(io.BytesIO(download.content))
        self.assertEqual([p.text for p in updated.paragraphs], ["KEEP: original first paragraph.", "Updated paragraph."])
        preview = self.client.get(f"/api/files/{artifact['id']}/preview", headers={"X-Session-ID": "session-alice"})
        renderer_available = (Path(r"C:\Program Files\LibreOffice\program\soffice.exe").is_file()
                              or shutil.which("soffice") or shutil.which("soffice.exe"))
        if renderer_available:
            self.assertEqual(preview.status_code, 200, preview.text[:200] if not preview.content.startswith(b"%PDF") else "")
            self.assertEqual(preview.headers["content-type"], "application/pdf")
        else:
            self.assertEqual(preview.status_code, 422)
            self.assertIn("下載原檔", preview.json()["detail"])
        for route in ("preview", "download"):
            denied = self.client.get(f"/api/files/{artifact['id']}/{route}", headers={"X-Session-ID": "session-bob"})
            self.assertEqual(denied.status_code, 403)
        followup = [("read_file", {"file_id": artifact["id"]}),
                    ("edit_file", {"file_id": artifact["id"], "filename": "lesson-v3.docx",
                                   "operations": [{"op": "set_paragraph", "index": 1, "text": "Follow-up paragraph."}]})]
        with self.model(followup):
            more = self.events(self.ask("继续修改上一版 Word 第2段为 Follow-up paragraph."))
        self.assertEqual(more[-1]["type"], "done", more)
        next_artifact = more[-1]["artifacts"][0]
        self.assertEqual(next_artifact["parent_id"], artifact["id"])
        self.assertEqual(next_artifact["version"], artifact["version"] + 1)
        self.assertEqual(Document(source).paragraphs[1].text, "Change this paragraph.")
        reloaded = main.DataStore()
        self.assertEqual(reloaded._conversations["conversation-alice"]["messages"][-1]["artifacts"], more[-1]["artifacts"])

    def test_global_admission_and_cross_account_result_isolation(self):
        async def complete(agent, messages, tools):
            if not agent.runtime.generated:
                return {"content": "", "tool_calls": [{"id": "create", "type": "function", "function": {
                    "name": "create_file", "arguments": json.dumps({"format": "docx", "filename": "private.docx",
                        "content": {"paragraphs": ["PRIVATE_" + agent.runtime.user_id]}})}}]}
            return {"content": json.dumps({"status": "complete", "summary": "done", "unmet_requirements": [],
                                           "output_file_ids": list(agent.runtime.generated)})}

        async def run_both():
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://local.test") as client:
                return await asyncio.gather(*[client.post("/api/chat", headers={"X-Session-ID": "session-" + user},
                    json={"question": "生成个人 Word", "session_id": "session-" + user}) for user in ("alice", "bob")])

        with patch.object(DeepSeekFileAgent, "complete", complete):
            responses = asyncio.run(run_both())
            self.assertEqual(sorted(item.status_code for item in responses), [200, 429])
            # Once the admitted job ends, the rejected user can retry normally.
            for index, user in enumerate(("alice", "bob")):
                if responses[index].status_code == 429:
                    responses[index] = self.ask("生成个人 Word", user=user)
        for user, response in zip(("alice", "bob"), responses):
            events = self.events(response)
            self.assertEqual(events[-1]["type"], "done", events)
            artifact = events[-1]["artifacts"][0]
            good = self.client.get(f"/api/files/{artifact['id']}/download", headers={"X-Session-ID": "session-" + user})
            self.assertEqual(Document(io.BytesIO(good.content)).paragraphs[0].text, "PRIVATE_" + user)
            other = "bob" if user == "alice" else "alice"
            for operation in ("download", "preview"):
                bad = self.client.get(f"/api/files/{artifact['id']}/{operation}", headers={"X-Session-ID": "session-" + other})
                self.assertEqual(bad.status_code, 403)
            wrong = self.client.get(f"/api/conversation/conversation-{user}", headers={"X-Session-ID": "session-" + other})
            self.assertEqual(wrong.status_code, 403)

    def test_same_conversation_race_is_rejected_after_async_reference_lookup(self):
        original = self.store.get_all_files
        async def exercise():
            ready = asyncio.Event()
            entered = 0
            async def synchronized(user):
                nonlocal entered
                result = await original(user)
                entered += 1
                if entered == 2:
                    ready.set()
                await ready.wait()
                return result
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://local.test") as client:
                with patch.object(self.store, "get_all_files", synchronized):
                    return await asyncio.gather(*[client.post("/api/chat", headers={"X-Session-ID": "session-alice"},
                        json={"question": "生成 Word", "session_id": "session-alice"}) for _ in range(2)])
        with self.model([("create_file", {"format": "docx", "filename": "test.docx", "content": {"paragraphs": ["test"]}})]):
            responses = asyncio.run(exercise())
        self.assertEqual(sorted(item.status_code for item in responses), [200, 409])
        successful = next(item for item in responses if item.status_code == 200)
        self.assertEqual(self.events(successful)[-1]["type"], "done")

    def test_large_excel_generation_does_not_fail_due_to_searchable_preview(self):
        rows = [[f"entry-{r}-{c}" for c in range(10)] for r in range(120)]
        with self.model([("create_file", {"format": "xlsx", "filename": "large.xlsx", "content": {
                "sheets": [{"name": "Data", "headers": [f"H{c}" for c in range(10)], "rows": rows}]}})]):
            events = self.events(self.ask("生成120行10列Excel"))
        self.assertEqual(events[-1]["type"], "done", events)
        artifact = events[-1]["artifacts"][0]
        raw = self.client.get(f"/api/files/{artifact['id']}/download", headers={"X-Session-ID": "session-alice"})
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw.content))
        self.assertEqual(workbook.active.max_row, 121)
        self.assertEqual(workbook.active["J121"].value, "entry-119-9")


if __name__ == "__main__":
    unittest.main()
