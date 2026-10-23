"""Offline file-agent boundary tests with real, isolated MCP stdio workers.

Fixtures live in temporary directories. DeepSeek HTTP is handled by an in-memory
transport; document tools and their subprocesses execute the production code.
"""
from __future__ import annotations

import asyncio
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from contextlib import asynccontextmanager
from pathlib import Path
from unittest.mock import AsyncMock, patch

import httpx
from docx import Document

import tests  # Establish synthetic-only storage before application imports.
from services.file_agent import agent as agent_module
from services.file_agent import runtime as runtime_module
from services.file_agent.agent import DeepSeekFileAgent, should_handle
from services.file_agent.runtime import FileAgentError, TaskRuntime, safe_child, user_key


class FileAgentRoutingTests(unittest.TestCase):
    def test_normal_questions_do_not_activate_file_tools(self):
        for question in ("你好", "解釋牛頓第二定律", "PDF 是甚麼格式？", "Excel 的 SUM 函數是甚麼？"):
            with self.subTest(question=question):
                self.assertFalse(should_handle(question))

    def test_referenced_file_questions_without_edit_request_stay_in_normal_chat(self):
        for question in ("請解釋資料的主要觀點", "根據校曆，今天有甚麼活動？"):
            with self.subTest(question=question):
                self.assertFalse(should_handle(question, "校曆.pdf", ["source-id"]))

    def test_creation_and_followup_edit_requests_activate_tools(self):
        self.assertTrue(should_handle("生成一份 PDF"))
        self.assertTrue(should_handle("把資料整理成 Excel", "校曆.pdf", ["source-id"]))
        self.assertTrue(should_handle("修改第二段", messages=[{"is_user": False, "artifacts": [{"id": "generated"}]}]))
        self.assertFalse(should_handle("繼續解釋這個概念", messages=[{"is_user": False, "artifacts": [{"id": "generated"}]}]))


class FileAgentRuntimeTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="hk_file_agent_runtime_")
        self.base = Path(self.temp.name)
        self.runtime = TaskRuntime(self.base / "uploads", "alice", "conversation-a")

    def tearDown(self):
        self.temp.cleanup()

    def add_word(self, runtime=None, file_id="source-a", text="ALICE_PRIVATE_DOCUMENT"):
        runtime = runtime or self.runtime
        source = self.base / (file_id + ".docx")
        doc = Document()
        doc.add_paragraph(text)
        doc.save(source)
        runtime.add_input({"id": file_id, "name": source.name}, source)
        return source

    def test_task_roots_are_account_and_task_scoped(self):
        same_account = TaskRuntime(self.base / "uploads", "alice", "conversation-b")
        other_account = TaskRuntime(self.base / "uploads", "bob", "conversation-a")
        self.assertNotEqual(self.runtime.root, same_account.root)
        self.assertNotEqual(self.runtime.root.parent, other_account.root.parent)
        self.assertEqual(self.runtime.root.parent.name, user_key("alice"))
        state = json.loads((self.runtime.root / "task.json").read_text(encoding="utf-8"))
        self.assertEqual(state["task_id"], self.runtime.task_id)
        self.assertEqual(state["conversation_id"], "conversation-a")
        self.assertNotIn("alice", state["user_key"])

    def test_path_traversal_and_sibling_prefix_are_rejected(self):
        for candidate in (self.runtime.outputs / ".." / "escaped.docx",
                          self.runtime.outputs.with_name("outputs-other") / "escaped.docx",
                          self.base / "escaped.docx"):
            with self.subTest(candidate=candidate):
                with self.assertRaises(FileAgentError):
                    safe_child(self.runtime.outputs, candidate)
        self.assertEqual(safe_child(self.runtime.outputs, self.runtime.outputs / "safe.docx"),
                         (self.runtime.outputs / "safe.docx").resolve())

    def test_input_is_a_copy_and_size_quota_is_enforced(self):
        source = self.add_word()
        copied = self.runtime.paths["source-a"]
        self.assertNotEqual(copied, source)
        self.assertTrue(copied.is_relative_to(self.runtime.inputs))
        source.write_bytes(b"changed outside the task")
        self.assertEqual(Document(copied).paragraphs[0].text, "ALICE_PRIVATE_DOCUMENT")
        with patch.object(runtime_module, "MAX_INPUT_BYTES", 1):
            with self.assertRaisesRegex(FileAgentError, "50 MB"):
                self.runtime.add_input({"id": "too-big"}, source)
        self.assertNotIn("too-big", self.runtime.paths)

    def test_task_storage_quota_raises_instead_of_reporting_success(self):
        with patch.object(runtime_module, "MAX_TASK_BYTES", 1):
            with self.assertRaisesRegex(FileAgentError, "150 MB"):
                self.runtime.check_quota()
        self.assertFalse(self.runtime.generated)

    async def test_real_stdio_read_returns_document_and_hides_local_paths(self):
        self.add_word()
        result = await self.runtime.call("read_file", {"file_id": "source-a"})
        serialized = json.dumps(result, ensure_ascii=False)
        self.assertIn("ALICE_PRIVATE_DOCUMENT", serialized)
        self.assertNotIn(str(self.runtime.root), serialized)
        self.assertIn("source-a", self.runtime.read_ids)
        self.assertTrue((self.runtime.root / "mcp-context.json").is_file())

    async def test_unknown_ids_and_traversal_ids_are_rejected_by_real_worker(self):
        self.add_word()
        for file_id in ("missing-id", "../source-a", "../../../../config/settings.py"):
            with self.subTest(file_id=file_id):
                with self.assertRaises(FileAgentError):
                    await self.runtime.call("read_file", {"file_id": file_id})
        self.assertFalse(self.runtime.read_ids)
        self.assertFalse(self.runtime.generated)

    async def test_pagination_cannot_skip_unread_middle(self):
        source = self.base / "pages.docx"
        doc = Document()
        for text in ("ONE", "TWO", "THREE"):
            doc.add_paragraph(text)
        doc.save(source)
        self.runtime.add_input({"id": "paged", "name": source.name}, source)
        await self.runtime.call("read_file", {"file_id": "paged", "start": 0, "limit": 1})
        self.assertNotIn("paged", self.runtime.read_ids)
        await self.runtime.call("read_file", {"file_id": "paged", "start": 2, "limit": 1})
        self.assertNotIn("paged", self.runtime.read_ids)
        await self.runtime.call("read_file", {"file_id": "paged", "start": 1, "limit": 1})
        self.assertIn("paged", self.runtime.read_ids)

    async def test_image_without_ocr_is_only_an_asset_not_readable_text(self):
        import fitz
        source = self.base / "image.png"
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
        pix.clear_with(255)
        pix.save(source)
        self.runtime.add_input({"id": "image", "name": source.name, "extraction_status": "failed"}, source)
        result = await self.runtime.call("read_file", {"file_id": "image"})
        self.assertFalse(result["capabilities"]["text_readable"])
        self.assertNotIn("image", self.runtime.read_ids)
        self.assertIn("image", self.runtime.asset_ids)

    async def test_scanned_pdf_without_complete_ocr_is_an_explicit_error(self):
        import fitz
        source = self.base / "scan.pdf"
        pdf = fitz.open()
        pdf.new_page()
        pdf.save(source)
        pdf.close()
        self.runtime.add_input({"id": "scan", "name": source.name, "extraction_status": "partial"},
                               source, "--- 第2页 ---\nPARTIAL OCR")
        with self.assertRaisesRegex(FileAgentError, "OCR"):
            await self.runtime.call("read_file", {"file_id": "scan"})
        self.assertNotIn("scan", self.runtime.read_ids)

    async def test_unknown_operation_is_rejected_before_worker_launch(self):
        with self.assertRaises(FileAgentError):
            await self.runtime.call("run_python", {"code": "print('not executable')"})
        self.assertFalse((self.runtime.root / "mcp-context.json").exists())

    async def test_real_worker_receives_no_provider_secrets_and_uses_task_temp(self):
        from mcp.client.stdio import stdio_client as original_stdio
        self.add_word()
        captured = []

        @asynccontextmanager
        async def checked(parameters, **kwargs):
            captured.append(parameters.env)
            async with original_stdio(parameters, **kwargs) as streams:
                yield streams

        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": "DO_NOT_INHERIT_PROVIDER_SECRET",
                                     "PRIVATE_TASK_TEST_SECRET": "DO_NOT_INHERIT_OTHER_SECRET"}), \
             patch("mcp.client.stdio.stdio_client", checked):
            result = await self.runtime.call("read_file", {"file_id": "source-a"})
        self.assertIn("ALICE_PRIVATE_DOCUMENT", json.dumps(result))
        self.assertEqual(len(captured), 1)
        self.assertNotIn("DEEPSEEK_API_KEY", captured[0])
        self.assertNotIn("PRIVATE_TASK_TEST_SECRET", captured[0])
        self.assertEqual(captured[0]["TEMP"], str(self.runtime.root))
        self.assertEqual(captured[0]["TMP"], str(self.runtime.root))

    async def test_two_accounts_read_concurrently_without_cross_account_ids(self):
        self.add_word()
        bob = TaskRuntime(self.base / "uploads", "bob", "conversation-b")
        self.add_word(bob, "source-b", "BOB_PRIVATE_DOCUMENT")
        a, b = await asyncio.gather(
            self.runtime.call("read_file", {"file_id": "source-a"}),
            bob.call("read_file", {"file_id": "source-b"}),
        )
        self.assertIn("ALICE_PRIVATE_DOCUMENT", json.dumps(a))
        self.assertNotIn("BOB_PRIVATE_DOCUMENT", json.dumps(a))
        self.assertIn("BOB_PRIVATE_DOCUMENT", json.dumps(b))
        self.assertNotIn("ALICE_PRIVATE_DOCUMENT", json.dumps(b))
        with self.assertRaises(FileAgentError):
            await bob.call("read_file", {"file_id": "source-a"})
        with self.assertRaises(FileAgentError):
            await self.runtime.call("read_file", {"file_id": "source-b"})

    async def test_same_account_other_task_file_is_rejected_at_startup(self):
        other = TaskRuntime(self.base / "uploads", "alice", "conversation-a")
        self.add_word(other, "other-task", "OTHER_TASK_SECRET")
        self.runtime.paths["forged-id"] = other.paths["other-task"]
        with self.assertRaises(FileAgentError) as caught:
            await self.runtime.call("read_file", {"file_id": "forged-id"})
        self.assertNotIn("OTHER_TASK_SECRET", str(caught.exception))
        self.assertFalse(self.runtime.generated)

    async def test_corrupted_office_file_returns_error_without_success(self):
        source = self.base / "broken.docx"
        source.write_bytes(b"not a zip or a valid Office document")
        self.runtime.add_input({"id": "broken", "name": "broken.docx"}, source)
        with self.assertRaises(FileAgentError):
            await self.runtime.call("read_file", {"file_id": "broken"})
        self.assertFalse(self.runtime.generated)
        self.assertNotIn("broken", self.runtime.read_ids)

    async def test_worker_timeout_is_an_error_and_has_no_generated_artifact(self):
        self.add_word()
        async def slow_call(*args, **kwargs):
            await asyncio.sleep(8)
        with patch.object(runtime_module, "TOOL_TIMEOUT", 3), patch("mcp.ClientSession.call_tool", slow_call):
            with self.assertRaisesRegex(FileAgentError, "逾時|無法執行"):
                await self.runtime.call("read_file", {"file_id": "source-a"})
        self.assertFalse(self.runtime.generated)
        self.assertFalse(list(self.runtime.outputs.iterdir()))

    def test_actual_worker_memory_limit_is_set_on_windows(self):
        if os.name != "nt":
            self.skipTest("Windows Job Object check")
        script = (
            "from services.file_agent.mcp_server import _resource_limit; "
            "import win32job; j=_resource_limit(); "
            "d=win32job.QueryInformationJobObject(j,win32job.JobObjectExtendedLimitInformation); "
            "print(d['ProcessMemoryLimit']); "
            "print(bool(d['BasicLimitInformation']['LimitFlags'] & win32job.JOB_OBJECT_LIMIT_PROCESS_MEMORY))"
        )
        child = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True,
                               timeout=15, cwd=str(Path(__file__).resolve().parents[1]),
                               creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        self.assertEqual(child.returncode, 0, child.stderr)
        self.assertEqual(child.stdout.splitlines(), [str(768 * 1024**2), "True"])


class DeepSeekFileAgentHttpTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="hk_file_agent_http_")
        self.runtime = TaskRuntime(Path(self.temp.name) / "uploads", "alice", "conversation-a")
        self.config = {"DEEPSEEK_API_KEY": "offline-test-key", "DEEPSEEK_API_URL": "https://deepseek.invalid/chat/completions"}
        self.authorize = AsyncMock()
        self.agent = DeepSeekFileAgent(self.runtime, self.config, self.authorize)
        self.requests = []
        self.real_client = httpx.AsyncClient

    def tearDown(self):
        self.temp.cleanup()

    def transport(self, handler):
        def factory(*args, **kwargs):
            return self.real_client(*args, **kwargs, transport=httpx.MockTransport(handler))
        return patch.object(agent_module.httpx, "AsyncClient", side_effect=factory)

    def response(self, request, message, reason="stop"):
        body = json.loads(request.content)
        self.requests.append(body)
        self.assertEqual(request.headers["authorization"], "Bearer offline-test-key")
        self.assertEqual(body["model"], "deepseek-flash")
        self.assertEqual(body["thinking"], {"type": "disabled"})
        self.assertFalse(body["stream"])
        self.assertEqual(body["tool_choice"], "auto")
        return httpx.Response(200, json={"choices": [{"finish_reason": reason, "message": message}]}, request=request)

    async def test_text_only_claim_without_actual_file_is_rejected(self):
        with self.transport(lambda req: self.response(req, {"role": "assistant", "content": "已完成文件"})):
            with self.assertRaises(FileAgentError):
                [event async for event in self.agent.run("生成 Word", [])]
        self.assertFalse(self.runtime.generated)

    async def test_truncated_provider_reply_is_rejected(self):
        with self.transport(lambda req: self.response(req, {"role": "assistant", "content": "unfinished"}, "length")):
            with self.assertRaisesRegex(FileAgentError, "不完整"):
                await self.agent.complete([], [])

    async def test_http_error_is_reported_without_provider_body_or_secret(self):
        def failed(req):
            return httpx.Response(429, text="PRIVATE_PROVIDER_BODY offline-test-key", request=req)
        with self.transport(failed):
            with self.assertRaisesRegex(FileAgentError, "HTTP 429") as caught:
                await self.agent.complete([], [])
        self.assertNotIn("PRIVATE_PROVIDER_BODY", str(caught.exception))
        self.assertNotIn("offline-test-key", str(caught.exception))

    async def test_authorization_is_checked_before_any_http_request(self):
        self.authorize.side_effect = FileAgentError("permission revoked")
        with self.transport(lambda req: self.fail("unauthorized request reached HTTP transport")):
            with self.assertRaisesRegex(FileAgentError, "permission revoked"):
                [event async for event in self.agent.run("生成 PDF", [])]

    async def test_task_timeout_never_emits_success(self):
        async def delayed(req):
            await asyncio.sleep(0.2)
            return self.response(req, {"role": "assistant", "content": "已完成"})
        events = []
        with self.transport(delayed), patch.object(agent_module, "TASK_TIMEOUT", 0.01):
            with self.assertRaises(TimeoutError):
                async for event in self.agent.run("生成 Word", []):
                    events.append(event)
        self.assertFalse(any(event["type"] == "result" for event in events))
        self.assertFalse(self.runtime.generated)


# Exercise the existing /api/chat and legacy /api/ask handlers using their
# temporary-store fixture, with all search/model network calls disabled.
_fixture_path = Path(__file__).with_name("test_web_search_api.py")
_fixture_spec = importlib.util.spec_from_file_location("_file_agent_web_fixture", _fixture_path)
_fixture_module = importlib.util.module_from_spec(_fixture_spec)
_fixture_spec.loader.exec_module(_fixture_module)


class FileAgentNormalChatApiTests(_fixture_module.WebSearchApiFixture):
    def test_existing_and_legacy_routes_preserve_normal_chat(self):
        for route in self.routes:
            with self.subTest(route=route):
                self.reset_conversation()
                _, events = self.ask(route, question="解釋牛頓第二定律", search=False)
                self.assertEqual(events[-1]["type"], "done")
                self.assertTrue(any(event["type"] == "content" for event in events))
                self.assertFalse(any(event["type"] in {"file_status", "artifacts"} for event in events))
                self.assertEqual(self.latest_assistant()["content"], self.answer)

    def test_explanatory_question_with_reference_does_not_start_file_task(self):
        for route in self.routes:
            with self.subTest(route=route):
                self.reset_conversation()
                _, events = self.ask(route, question="根據這份資料解釋主要內容", search=False,
                                     file_references="fixture.txt")
                self.assertEqual(events[-1]["type"], "done")
                self.assertFalse(any(event["type"] in {"file_status", "artifacts"} for event in events))


if __name__ == "__main__":
    unittest.main()
