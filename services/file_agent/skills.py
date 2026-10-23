"""Local, capability-based Office/PDF skills. Only trusted file IDs address files.

create_file: {format, filename, content}. content uses paragraphs/tables for DOCX
and PDF, slides:[{title,body,table}] for PPTX, sheets:[{name,headers,rows}]
for XLSX. edit_file accepts {file_id,filename,operations:[{op,...}]}.
Indexes are zero based except spreadsheet A1 cell coordinates.
"""
from __future__ import annotations

import copy
import csv
import io
import json
import math
import re
import uuid
import zipfile
from pathlib import Path
from typing import Any

import pymupdf as fitz
from docx import Document
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


class FileSkillError(ValueError):
    """Safe, user-readable file-operation failure."""


_FORMATS = {"docx", "pptx", "xlsx", "pdf"}
_MAX_BYTES = 50 * 1024 * 1024
_DANGEROUS_FORMULA = re.compile(r"(?:WEBSERVICE|HYPERLINK|RTD|DDE|CALL|REGISTER|EXEC|IMPORTXML|IMPORTDATA|IMAGE)\s*\(|\[[^\]]+\]|https?://|file://|\\\\|\|", re.I)

_STYLE_FIELDS = "font_size font_name bold color".split()
_EDIT_FIELDS = {
    "docx": {
        "replace_text": "old new", "set_paragraph": "index text", "set_table_cell": "table row col text",
        "add_paragraph": "text style", "add_table": "headers rows", "add_image": "file_id width_inches",
        "replace_image": "index file_id", "remove_image": "index", "add_equation": "text",
        "style_paragraph": "index " + " ".join(_STYLE_FIELDS),
    },
    "pptx": {
        "replace_text": "old new", "set_shape_text": "slide shape text", "set_table_cell": "slide shape row col text",
        "add_slide": "title body table image notes background title_color text_color",
        "style_shape": "slide shape " + " ".join(_STYLE_FIELDS),
        "add_image": "slide file_id x y width height", "replace_image": "slide shape file_id",
    },
    "xlsx": {
        "set_cell": "sheet cell value", "set_range": "sheet start values",
        "style_range": "sheet range font_size font_name bold color fill number_format wrap_text",
        "merge_cells": "sheet range", "column_width": "sheet column width",
        "add_image": "sheet file_id cell width_px height_px", "replace_image": "sheet index file_id",
        "remove_image": "sheet index",
    },
    "pdf": {"replace_text": "page old new font_size", "add_text": "page text x y right bottom font_size",
            "delete_pages": "pages", "add_image": "page file_id rect"},
}
# Only prefer fonts with CJK coverage; otherwise use PyMuPDF's bundled china-s.
_PDF_FONT = next((p for p in (Path("C:/Windows/Fonts/msyh.ttc"),
                               Path("C:/Windows/Fonts/simsun.ttc")) if p.is_file()), None)


def _pdf_font(page, unique=False):
    if _PDF_FONT:
        name = "agentedit" + uuid.uuid4().hex[:8] if unique else "agentfont"
        page.insert_font(fontname=name, fontfile=str(_PDF_FONT))
        return name, fitz.Font(fontfile=str(_PDF_FONT))
    return "china-s", fitz.Font(fontname="china-s")


def _schema(name, description, properties, required):
    return {"type": "function", "function": {"name": name, "description": description,
        "parameters": {"type": "object", "properties": properties,
                       "required": required, "additionalProperties": False}}}


TOOL_DEFINITIONS = [
    _schema("inspect_file", "查看已授权文件的类型、页/段/表/幻灯片信息。索引从0开始，Excel使用A1坐标。", {"file_id": {"type": "string"}}, ["file_id"]),
    _schema("read_file", "分页读取文件所有结构化内容和原文；有has_more时继续读next_start。DOCX单位为段落或表格，PPTX为幻灯片，PDF为页，XLSX为数据行。", {"file_id": {"type": "string"}, "start": {"type": "integer", "minimum": 0}, "limit": {"type": "integer", "minimum": 1, "maximum": 100}}, ["file_id"]),
    _schema("create_file", "真正创建文件。DOCX/PDF content={title,paragraphs:[字符串或{text,style}],tables:[{headers:[],rows:[[]]}]}；PPTX content={slides:[{title,body:[字符串],table:{headers:[],rows:[[]]}}]}；XLSX content={sheets:[{name,headers:[],rows:[[]],widths:[数字]}]}。文本必须完整，图片仅可引用已有image.file_id。", {"format": {"type": "string", "enum": sorted(_FORMATS)}, "filename": {"type": "string"}, "content": {"type": "object"}}, ["format", "filename", "content"]),
    _schema("edit_file", "保留未改内容另存新版本。operations每项必须op。DOCX:set_paragraph(index,text),replace_text(old,new),set_table_cell(table,row,col,text),add_paragraph(text),add_table(headers,rows),style_paragraph(index,font_size,bold,color)。PPTX:set_shape_text(slide,shape,text),replace_text(old,new),set_table_cell(slide,shape,row,col,text),add_slide(title,body),style_shape(slide,shape,font_size,bold,color)。XLSX:set_cell(sheet,cell,value),set_range(sheet,start,values),style_range(sheet,range,font_size,bold,color,fill,number_format),merge_cells(sheet,range),column_width(sheet,column,width)。PDF:replace_text(page,old,new),add_text(page,text,x,y,font_size),delete_pages(pages)。", {"file_id": {"type": "string"}, "filename": {"type": "string"}, "operations": {"type": "array", "items": {"type": "object"}, "minItems": 1, "maxItems": 200}}, ["file_id", "filename", "operations"]),
    _schema("validate_file", "重新打开并检查真实文件结构、大小、外链和危险公式。", {"file_id": {"type": "string"}}, ["file_id"]),
]


def tool_definitions():
    return copy.deepcopy(TOOL_DEFINITIONS)


TOOL_DEFINITIONS[3]["function"]["description"] += " 图片修改：DOCX replace_image(index,file_id),remove_image(index),add_image(file_id,width_inches)；PPTX replace_image(slide,shape,file_id),add_image(slide,file_id,x,y,width,height)，坐标单位英寸；XLSX add_image(sheet,file_id,cell,width_px,height_px),replace_image(sheet,index,file_id),remove_image(sheet,index)。DOCX add_equation(text)添加可编辑Office线性公式。所有图片引用已授权file_id。"
TOOL_DEFINITIONS[3]["function"]["description"] += " DOCX style_paragraph、PPTX style_shape、XLSX style_range 均可指定font_name字体名称。不支持的参数会明确拒绝，不会忽略。"


def _text(value):
    result = str(value if value is not None else "")
    if len(result) > 250000 or any(ord(c) < 32 and c not in "\n\r\t" for c in result):
        raise FileSkillError("文字过长或含不支持的控制字符")
    return result


def _index(value, length, label="索引"):
    if isinstance(value, bool) or not isinstance(value, int) or not 0 <= value < length:
        raise FileSkillError(f"{label}不存在，请先读取文件确认位置")
    return value


def _color(value):
    value = str(value).lstrip("#")
    if not re.fullmatch(r"[0-9a-fA-F]{6}", value):
        raise FileSkillError("颜色须为六位十六进制，例如 #1F4E78")
    return value.upper()


def _value(value):
    if value is not None and not isinstance(value, (str, int, float, bool)):
        raise FileSkillError("单元格值须为文字、数字或布尔值")
    if isinstance(value, float) and not math.isfinite(value):
        raise FileSkillError("单元格数字必须有限")
    if isinstance(value, str):
        _text(value)
        if value.startswith("=") and _DANGEROUS_FORMULA.search(value):
            raise FileSkillError("拒绝包含外部连接或系统调用的公式")
        if len(value) > 32767:
            raise FileSkillError("Excel单元格文字超过32767字符")
    return value


def _replace_paragraph(paragraph, old, new):
    """Replace across runs while retaining untouched run properties and drawings."""
    if not old:
        raise FileSkillError("替换原文不能为空")
    count = 0
    # Work from the end so replacement text containing old does not get replaced again.
    joined = "".join(r.text for r in paragraph.runs)
    positions = [m.start() for m in re.finditer(re.escape(old), joined)]
    for start in reversed(positions):
        runs = list(paragraph.runs)
        end = start + len(old)
        offset = 0
        inserted = False
        for run in runs:
            begin, finish = offset, offset + len(run.text)
            offset = finish
            if finish <= start or begin >= end:
                continue
            left = run.text[:max(0, start - begin)]
            right = run.text[max(0, end - begin):] if end < finish else ""
            replacement = left + (new if not inserted else "") + right
            # python-docx Run.text clears drawings; modify only text-bearing children.
            if run._r.tag == qn("w:r"):
                for node in list(run._r):
                    if node.tag in {qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")}:
                        run._r.remove(node)
                run._r.add_t(replacement)
            else:
                run.text = replacement
            inserted = True
        count += int(inserted)
    return count


def _docx_paragraphs(doc):
    yield from doc.paragraphs
    def table_paragraphs(table):
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                yield from cell.paragraphs
                for nested in cell.tables:
                    yield from table_paragraphs(nested)
    for table in doc.tables:
        yield from table_paragraphs(table)
    for section in doc.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                yield from table_paragraphs(table)


class FileSkills:
    def __init__(self, input_paths: dict[str, Path], output_dir: Path, **_kwargs):
        self.input_paths = {str(k): Path(v).resolve() for k, v in input_paths.items()}
        self.output_dir = Path(output_dir).resolve()
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def execute(self, operation: str, arguments: dict) -> dict:
        if operation not in {"inspect_file", "read_file", "create_file", "edit_file", "validate_file"}:
            raise FileSkillError("不支持的文件工具")
        if not isinstance(arguments, dict):
            raise FileSkillError("工具参数必须是JSON对象")
        try:
            return getattr(self, operation)(**arguments)
        except FileSkillError:
            raise
        except Exception as exc:
            raise FileSkillError(f"文件操作失败（{type(exc).__name__}），请检查格式和目标位置") from exc

    def _source(self, file_id):
        path = self.input_paths.get(str(file_id))
        if path is None or not path.is_file():
            raise FileSkillError("文件不存在或当前任务无权访问")
        self._guard(path)
        return path

    def _guard(self, path):
        if path.stat().st_size > _MAX_BYTES:
            raise FileSkillError("文件超过50MB处理限制")
        if path.suffix.lower().lstrip(".") in {"docx", "pptx", "xlsx"}:
            try:
                from services.resource_limits import validate_zip, bounded_member
                validate_zip(path)
                with zipfile.ZipFile(path) as archive:
                    items = archive.infolist()
                    for item in items:
                        name = item.filename.lower()
                        if name.endswith("vbaproject.bin") or "/embeddings/" in name or name.startswith("xl/externallinks/") or name == "xl/connections.xml" or name.startswith("xl/querytables/"):
                            raise FileSkillError("不支持宏、嵌入程序或外部工作簿连接")
                        if name.endswith((".xml", ".rels")):
                            xml_bytes = bounded_member(archive, item)
                            if re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", xml_bytes, re.I):
                                raise FileSkillError("Office文件包含不允许的XML实体或文档类型声明")
                            from lxml import etree
                            xml_root = etree.fromstring(xml_bytes, parser=etree.XMLParser(resolve_entities=False, no_network=True))
                            if name.startswith("word/"):
                                instructions = xml_root.xpath("//*[local-name()='instrText']/text()")
                                fields = " ".join(instructions) + " " + "".join(instructions)
                                fields += " " + " ".join(xml_root.xpath("//*[local-name()='fldSimple']/@*[local-name()='instr']"))
                                if re.search(r"\b(?:DDE(?:AUTO)?|INCLUDETEXT|INCLUDEPICTURE|LINK|DATABASE)\b", fields, re.I):
                                    raise FileSkillError("Word含DDE或外部资源字段，不能安全处理")
                                if xml_root.xpath("//*[local-name()='altChunk']"):
                                    raise FileSkillError("Word包含外来嵌入内容，无法保证安全保留")
                            if name.startswith("xl/"):
                                for formula in xml_root.xpath("//*[local-name()='f' or local-name()='definedName']/text()"):
                                    _value("=" + formula)
                        if name.endswith(".rels"):
                            for rel in xml_root:
                                # Ordinary clickable hyperlinks are inert; do not follow them.
                                if rel.get("TargetMode") == "External" and not rel.get("Type", "").endswith("/hyperlink"):
                                    raise FileSkillError("文件含外部资源连接，不能安全处理")
            except zipfile.BadZipFile as exc:
                raise FileSkillError("Office文件已损坏或格式不正确") from exc
            except ValueError as exc:
                raise FileSkillError(str(exc)) from exc
        elif path.suffix.lower() == ".pdf":
            try:
                with fitz.open(path) as pdf:
                    from services.resource_limits import validate_pdf
                    validate_pdf(pdf)
                    if pdf.needs_pass:
                        raise FileSkillError("PDF已加密，需要先解密")
                    if pdf.xref_length() > 100000:
                        raise FileSkillError("PDF内部对象过多")
                    for xref in range(1, pdf.xref_length()):
                        obj = pdf.xref_object(xref)
                        if re.search(r"/(?:JavaScript|JS|Launch|OpenAction|AA|EmbeddedFile)\b", obj):
                            raise FileSkillError("PDF包含脚本、自动动作或嵌入附件，不能安全处理")
            except FileSkillError:
                raise
            except Exception as exc:
                raise FileSkillError("PDF已损坏或无法读取") from exc

    def _target(self, filename, format):
        if format not in _FORMATS:
            raise FileSkillError("仅支持PDF、DOCX、PPTX、XLSX输出")
        name = _text(filename).strip()
        if not name or len(name) > 160 or re.search(r'[<>:"/\\|?*\x00-\x1f]', name) or name in {".", ".."} or name.endswith((".", " ")):
            raise FileSkillError("输出名称必须为普通文件名，不能包含路径")
        if Path(name).suffix.lower() != "." + format:
            raise FileSkillError("输出文件扩展名与格式不一致")
        path = self.output_dir / (uuid.uuid4().hex + "_" + name)
        if path.resolve().parent != self.output_dir:
            raise FileSkillError("输出路径无效")
        return path, name

    def _register(self, path, filename, summary):
        key = uuid.uuid4().hex
        self.input_paths[key] = path
        try:
            validation = self.validate_file(key)
        except Exception:
            self.input_paths.pop(key, None)
            path.unlink(missing_ok=True)
            raise
        return {"file_id": key, "filename": filename, "path": str(path), "format": path.suffix[1:],
                "validation": validation, "summary": summary}

    def inspect_file(self, file_id):
        path = self._source(file_id)
        ext = path.suffix.lower()[1:]
        info = {"file_id": file_id, "filename": path.name, "format": ext, "size": path.stat().st_size}
        if ext == "docx":
            doc = Document(path)
            info.update(paragraph_count=len(doc.paragraphs), table_count=len(doc.tables), image_count=len(doc.inline_shapes))
        elif ext == "pptx":
            prs = Presentation(path)
            info.update(slide_count=len(prs.slides), page_count=len(prs.slides))
        elif ext == "xlsx":
            wb = load_workbook(path, data_only=False, keep_links=False)
            try:
                info["sheets"] = [{"name": ws.title, "rows": ws.max_row, "columns": ws.max_column,
                                   "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
                                   "image_count": len(ws._images), "chart_count": len(ws._charts)} for ws in wb]
            finally:
                wb.close()
        elif ext == "pdf":
            with fitz.open(path) as pdf:
                if pdf.needs_pass:
                    raise FileSkillError("PDF已加密，需要先提供可读取的版本")
                info.update(page_count=len(pdf))
        elif ext in {"png", "jpg", "jpeg", "gif", "bmp", "webp"}:
            from PIL import Image
            with Image.open(path) as image:
                info.update(width=image.width, height=image.height, needs_ocr=True)
        elif ext not in {"txt", "md", "csv", "json"}:
            raise FileSkillError("暂不支持此文件格式")
        return info

    def read_file(self, file_id, start=0, limit=100):
        if isinstance(start, bool) or not isinstance(start, int) or start < 0 or not isinstance(limit, int) or not 1 <= limit <= 100:
            raise FileSkillError("分页参数无效")
        path = self._source(file_id)
        ext = path.suffix.lower()[1:]
        units = []
        unit = "item"
        if ext == "docx":
            doc = Document(path)
            unit = "paragraph_or_table"
            pi = ti = 0
            for child in doc.element.body:
                if child.tag == qn("w:p"):
                    paragraph = doc.paragraphs[pi]
                    units.append({"type": "paragraph", "index": pi, "text": paragraph.text, "style": paragraph.style.name})
                    pi += 1
                elif child.tag == qn("w:tbl"):
                    table = doc.tables[ti]
                    units.append({"type": "table", "index": ti, "rows": [[c.text for c in row.cells] for row in table.rows]})
                    ti += 1
            for idx, section in enumerate(doc.sections):
                for label, part in (("header", section.header), ("footer", section.footer)):
                    text = "\n".join(p.text for p in part.paragraphs)
                    if text.strip():
                        units.append({"type": label, "section": idx, "text": text})
            for idx, image in enumerate(doc.inline_shapes):
                units.append({"type": "image", "index": idx, "width_inches": image.width / 914400,
                              "height_inches": image.height / 914400, "message": "内嵌图片保留，替换时按此index指定"})
        elif ext == "pptx":
            prs = Presentation(path)
            unit = "slide"
            for si, slide in enumerate(prs.slides):
                shapes = []
                for idx, shape in enumerate(slide.shapes):
                    item = {"index": idx, "name": shape.name, "type": str(shape.shape_type)}
                    if shape.has_text_frame:
                        item["text"] = shape.text
                    if shape.has_table:
                        item["rows"] = [[c.text for c in row.cells] for row in shape.table.rows]
                    shapes.append(item)
                notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ""
                units.append({"type": "slide", "index": si, "shapes": shapes, "notes": notes})
        elif ext == "xlsx":
            wb = load_workbook(path, data_only=False, read_only=True, keep_links=False)
            unit = "sheet_row"
            try:
                for ws in wb:
                    if ws.max_row * ws.max_column > 300000:
                        raise FileSkillError("工作表过大，请先缩小处理范围")
                    for row_num, row in enumerate(ws.iter_rows(), 1):
                        units.append({"type": "row", "sheet": ws.title, "row": row_num,
                                      "cells": [{"cell": c.coordinate, "value": c.value.isoformat() if hasattr(c.value, "isoformat") else c.value} for c in row]})
            finally:
                wb.close()
        elif ext == "pdf":
            unit = "page"
            with fitz.open(path) as pdf:
                if pdf.needs_pass:
                    raise FileSkillError("PDF已加密，需要先解密")
                for index, page in enumerate(pdf):
                    units.append({"type": "page", "index": index, "text": page.get_text("text"), "width": page.rect.width, "height": page.rect.height,
                                  "needs_ocr": not bool(page.get_text("text").strip())})
        elif ext in {"txt", "md", "csv", "json"}:
            unit = "line"
            try:
                text = path.read_text(encoding="utf-8-sig")
            except UnicodeDecodeError:
                text = path.read_text(encoding="gb18030")
            units = [{"type": "line", "index": i, "text": line} for i, line in enumerate(text.splitlines())]
        elif ext in {"png", "jpg", "jpeg", "gif", "bmp", "webp"}:
            info = self.inspect_file(file_id)
            units = [{"type": "image", "index": 0, "width": info["width"], "height": info["height"],
                      "needs_ocr": True, "text": "", "message": "图片内容需使用服务器既有OCR结果；不能猜测图片中的文字"}]
        else:
            raise FileSkillError("暂不支持此文件格式")
        selected = units[start:start + limit]
        # Refuse oversized individual objects instead of silently dropping their content.
        if len(json.dumps(selected, ensure_ascii=False)) > 1800000:
            raise FileSkillError("此页内容过大，请减小limit分页读取")
        total = len(units)
        more = start + len(selected) < total
        paging = {"start": start, "returned": len(selected), "total": total, "has_more": more,
                  "next_start": start + len(selected) if more else None, "unit": unit}
        result = {"file_id": file_id, "format": ext, "items": selected, "content": selected,
                  "pagination": paging, "has_more": more, "next_start": paging["next_start"], "total": total, "truncated": False}
        if ext in {"txt", "md", "csv", "json"} and start == 0 and not more:
            # Only trusted local in-memory streams; no URL or plugin dispatch is enabled.
            from markitdown import MarkItDown
            result["markdown"] = MarkItDown(enable_plugins=False).convert_stream(
                io.BytesIO(text.encode("utf-8")), file_extension=path.suffix).text_content
            result["markdown_converter"] = "markitdown-local-no-plugins"
        return result

    def create_file(self, format, filename, content):
        format = str(format).lower().lstrip(".")
        path, filename = self._target(filename, format)
        if not isinstance(content, dict) or len(json.dumps(content, ensure_ascii=False)) > 2000000:
            raise FileSkillError("文件内容须为有效且大小受限的JSON对象")
        try:
            if format == "docx":
                doc = Document()
                normal = doc.styles["Normal"]
                normal.font.name = "Microsoft YaHei"
                normal.font.size = DocxPt(11)
                normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                if content.get("title"):
                    doc.add_heading(_text(content["title"]), 0)
                for item in content.get("paragraphs", []):
                    if isinstance(item, str):
                        doc.add_paragraph(_text(item))
                    else:
                        p = doc.add_paragraph(_text(item.get("text", "")), style=item.get("style"))
                        self._docx_style(p, item)
                for table in content.get("tables", []):
                    self._docx_table(doc, table)
                for image in content.get("images", []):
                    self._docx_image(doc, image)
                if not content.get("title") and not any(content.get(k) for k in ("paragraphs", "tables", "images")):
                    raise FileSkillError("Word内容不能为空")
                doc.save(path)
            elif format == "pptx":
                slides = content.get("slides", [])
                if not slides or len(slides) > 60:
                    raise FileSkillError("PPT必须包含1至60页完整幻灯片")
                prs = Presentation()
                prs.slide_width = Inches(13.333)
                prs.slide_height = Inches(7.5)
                for slide in slides:
                    self._pptx_slide(prs, slide)
                prs.save(path)
            elif format == "xlsx":
                sheets = content.get("sheets", [])
                if not sheets or len(sheets) > 30:
                    raise FileSkillError("Excel必须包含1至30个工作表")
                wb = Workbook()
                wb.remove(wb.active)
                names = set()
                for si, spec in enumerate(sheets):
                    name = _text(spec.get("name", f"Sheet{si + 1}"))
                    if len(name) > 31 or re.search(r'[\\/*?:\[\]]', name) or name.lower() in names:
                        raise FileSkillError("工作表名称无效或重复")
                    names.add(name.lower())
                    ws = wb.create_sheet(name)
                    headers = spec.get("headers", [])
                    rows = spec.get("rows", [])
                    if len(rows) > 10000 or any(len(row) > 200 for row in rows):
                        raise FileSkillError("工作表行列超过生成限制")
                    if headers:
                        ws.append([_value(v) for v in headers])
                    for row in rows:
                        ws.append([_value(v) for v in row])
                    if headers:
                        for c in ws[1]:
                            c.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
                            c.fill = PatternFill("solid", fgColor="1F4E78")
                        ws.freeze_panes = "A2"
                        ws.auto_filter.ref = ws.dimensions
                    for col in range(1, ws.max_column + 1):
                        widths = spec.get("widths", [])
                        width = widths[col - 1] if col <= len(widths) else min(55, max(14, max((len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 100) + 1)), default=0) + 4))
                        ws.column_dimensions[get_column_letter(col)].width = min(100, max(3, float(width)))
                    for row in ws:
                        for c in row:
                            c.alignment = Alignment(vertical="top", wrap_text=True)
                    for op in spec.get("formats", []):
                        self._xlsx_operation(wb, {"op": "style_range", "sheet": name, **op})
                try:
                    from openpyxl.workbook.properties import CalcProperties
                    wb.calculation = CalcProperties(fullCalcOnLoad=True)
                    wb.save(path)
                finally:
                    wb.close()
            else:
                self._create_pdf(path, content)
            return self._register(path, filename, "已创建并重新打开检查文件")
        except Exception:
            path.unlink(missing_ok=True)
            raise

    def _docx_style(self, p, spec):
        for run in p.runs:
            if "font_size" in spec:
                size = float(spec["font_size"])
                if not 6 <= size <= 96:
                    raise FileSkillError("字号超出支持范围")
                run.font.size = DocxPt(size)
            if "bold" in spec:
                run.bold = bool(spec["bold"])
            if "color" in spec:
                run.font.color.rgb = DocxColor.from_string(_color(spec["color"]))
            if "font_name" in spec:
                run.font.name = _text(spec["font_name"])
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _text(spec["font_name"]))

    def _docx_table(self, doc, spec):
        rows = ([spec["headers"]] if spec.get("headers") else []) + list(spec.get("rows", []))
        if not rows or len(rows) > 1000 or not max(map(len, rows)) or max(map(len, rows)) > 30:
            raise FileSkillError("Word表格须有内容且不超过1000行30列")
        table = doc.add_table(rows=len(rows), cols=max(map(len, rows)))
        table.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                table.cell(i, j).text = _text(value)
                if i == 0 and spec.get("headers"):
                    for run in table.cell(i, j).paragraphs[0].runs:
                        run.bold = True
        return table

    def _docx_image(self, doc, image):
        src, ratio = self._image_source(image["file_id"])
        width = float(image.get("width_inches", 5))
        if not 0.1 <= width <= 7:
            raise FileSkillError("图片宽度超出页面范围")
        height = width / ratio
        if height > 8.7:
            width, height = 8.7 * ratio, 8.7
        doc.add_picture(str(src), width=DocxInches(width), height=DocxInches(height))

    def _image_source(self, file_id):
        src = self._source(file_id)
        if src.suffix.lower() not in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}:
            raise FileSkillError("图片格式不支持")
        from PIL import Image
        with Image.open(src) as image:
            if image.width * image.height > 40000000:
                raise FileSkillError("图片像素超过4000万处理限制")
            ratio = image.width / image.height
            image.verify()
        return src, ratio

    def _pptx_slide(self, prs, spec):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(_color(spec.get("background", "F8FAFC")))
        title = slide.shapes.add_textbox(Inches(.65), Inches(.35), Inches(12.0), Inches(.8))
        title.text = _text(spec.get("title", ""))
        self._pptx_text_style(title, {"font_name": "Microsoft YaHei", "font_size": 30, "bold": True, "color": spec.get("title_color", "17365D")})
        self._pptx_check_text(title)
        body = spec.get("body", [])
        if isinstance(body, str):
            body = body.splitlines()
        if len(body) > 18 or any(len(_text(line)) > 350 for line in body):
            raise FileSkillError("幻灯片正文过多，请拆分为多页")
        has_table = bool(spec.get("table"))
        has_image = bool(spec.get("image"))
        height = 2.65 if has_table else 5.5
        width = 6.1 if has_image else 12.0
        shape = slide.shapes.add_textbox(Inches(.65), Inches(1.4), Inches(width), Inches(height))
        shape.text_frame.word_wrap = True
        shape.text = "\n".join(_text(line) for line in body)
        self._pptx_text_style(shape, {"font_name": "Microsoft YaHei", "font_size": min(24, max(14, 27 - len(body))), "color": spec.get("text_color", "263445")})
        self._pptx_check_text(shape)
        if has_table:
            table_spec = spec["table"]
            rows = ([table_spec["headers"]] if table_spec.get("headers") else []) + list(table_spec.get("rows", []))
            if not rows or len(rows) > 9 or max(map(len, rows)) > 8:
                raise FileSkillError("幻灯片表格最多9行8列，请拆分")
            cols = max(map(len, rows))
            table = slide.shapes.add_table(len(rows), cols, Inches(.65), Inches(4.15), Inches(12), Inches(2.7)).table
            for i, row in enumerate(rows):
                for j, value in enumerate(row):
                    table.cell(i, j).text = _text(value)
                    for p in table.cell(i, j).text_frame.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(min(14, 100 / len(rows)))
                            run.font.name = "Microsoft YaHei"
                    self._pptx_check_text(None, table.cell(i, j).text_frame, table.columns[j].width, table.rows[i].height)
        if has_image:
            image = spec["image"]
            src, ratio = self._image_source(image["file_id"])
            max_width, max_height = 5.5, 5.4
            width, height = min(max_width, max_height * ratio), min(max_height, max_width / ratio)
            slide.shapes.add_picture(str(src), Inches(7 + (max_width - width) / 2), Inches(1.5), width=Inches(width), height=Inches(height))
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = _text(spec["notes"])
        return slide

    def _pptx_text_style(self, shape, spec):
        if not shape.has_text_frame:
            raise FileSkillError("目标图形没有可编辑文字")
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                if "font_name" in spec:
                    run.font.name = _text(spec["font_name"])
                    properties = run._r.get_or_add_rPr()
                    east_asian = properties.find(qn("a:ea"))
                    if east_asian is None:
                        east_asian = OxmlElement("a:ea")
                        properties.append(east_asian)
                    east_asian.set("typeface", _text(spec["font_name"]))
                if "font_size" in spec:
                    size = float(spec["font_size"])
                    if not 6 <= size <= 96:
                        raise FileSkillError("字号超出范围")
                    run.font.size = Pt(size)
                if "bold" in spec:
                    run.font.bold = bool(spec["bold"])
                if "color" in spec:
                    run.font.color.rgb = RGBColor.from_string(_color(spec["color"]))

    def _pptx_check_text(self, shape, frame=None, width=None, height=None):
        """Conservative real-font fit check; refuse text that must overflow."""
        frame = frame if frame is not None else shape.text_frame
        width = float(width if width is not None else shape.width) / 12700
        height = float(height if height is not None else shape.height) / 12700
        width -= float(frame.margin_left + frame.margin_right) / 12700
        height -= float(frame.margin_top + frame.margin_bottom) / 12700
        if not frame.text.strip():
            return
        if width <= 0 or height <= 0:
            raise FileSkillError("幻灯片文字区域无可用空间")
        font = fitz.Font(fontfile=str(_PDF_FONT)) if _PDF_FONT else fitz.Font(fontname="china-s")
        needed = 0.0
        for p in frame.paragraphs:
            runs = list(p.runs)
            default_size = float(p.font.size.pt) if p.font.size is not None else 18.0
            line_width, line_height = 0.0, default_size * 1.2
            for run in runs:
                size = float(run.font.size.pt) if run.font.size is not None else default_size
                for char in run.text:
                    if char in "\n\v":
                        needed += line_height
                        line_width, line_height = 0.0, size * 1.2
                        continue
                    char_width = font.text_length(char, size)
                    if frame.word_wrap is not False and line_width and line_width + char_width > width:
                        needed += line_height
                        line_width, line_height = 0.0, size * 1.2
                    line_width += char_width
                    line_height = max(line_height, size * 1.2)
                    if frame.word_wrap is False and line_width > width + 2:
                        raise FileSkillError("幻灯片文字宽度超出图形，请缩短或拆分内容")
            needed += line_height
            needed += float(p.space_before.pt) if p.space_before is not None else 0
            needed += float(p.space_after.pt) if p.space_after is not None else 0
        if needed > height + 2:
            raise FileSkillError("幻灯片文字超出可用高度，请缩短、拆页或调整字号")

    def _create_pdf(self, path, content):
        pdf = fitz.open()
        font = fitz.Font(fontfile=str(_PDF_FONT)) if _PDF_FONT else fitz.Font(fontname="china-s")
        fontname = "agentfont" if _PDF_FONT else "china-s"
        page = None
        y = 0.0
        def new_page():
            nonlocal page, y
            if len(pdf) >= 100:
                raise FileSkillError("PDF超过100页限制")
            page = pdf.new_page(width=595, height=842)
            _pdf_font(page)
            y = 50.0
        def wrap(text, width, size):
            lines = []
            for para in _text(text).split("\n"):
                line = ""
                for char in para:
                    if line and font.text_length(line + char, size) > width:
                        lines.append(line)
                        line = char
                    else:
                        line += char
                lines.append(line)
            return lines
        def text(text, size=11, bold=False):
            nonlocal y
            for line in wrap(text, 495, size):
                if y + size * 1.6 > 790:
                    new_page()
                page.insert_text((50, y + size), line, fontname=fontname, fontsize=size, color=(.1, .16, .24))
                y += size * 1.6
            y += 8
        def table(spec):
            nonlocal y
            rows = ([spec["headers"]] if spec.get("headers") else []) + list(spec.get("rows", []))
            if not rows or len(rows) > 5000:
                raise FileSkillError("PDF表格为空或行数过多")
            cols = max(map(len, rows))
            if not 1 <= cols <= 12:
                raise FileSkillError("PDF表格列数超出范围")
            colw = 495 / cols
            for ri, row in enumerate(rows):
                lines = [wrap(value, colw - 12, 10) for value in list(row) + [""] * (cols - len(row))]
                height = max(len(x) for x in lines) * 15 + 12
                if height > 730:
                    raise FileSkillError("表格单行内容过多，请拆分")
                if y + height > 790:
                    new_page()
                for ci, cell_lines in enumerate(lines):
                    rect = fitz.Rect(50 + ci * colw, y, 50 + (ci + 1) * colw, y + height)
                    page.draw_rect(rect, color=(.65, .7, .75), fill=(.88, .93, .98) if ri == 0 and spec.get("headers") else None, width=.5)
                    for li, line in enumerate(cell_lines):
                        page.insert_text((rect.x0 + 6, rect.y0 + 16 + li * 15), line, fontname=fontname, fontsize=10)
                y += height
            y += 12
        try:
            pages = content.get("pages")
            if pages is None:
                pages = [content]
            if not pages or len(pages) > 100:
                raise FileSkillError("PDF页内容无效")
            for spec in pages:
                new_page()
                if spec.get("title"):
                    text(spec["title"], 20)
                for item in spec.get("paragraphs", []):
                    text(item if isinstance(item, str) else item.get("text", ""))
                for item in spec.get("tables", []):
                    table(item)
                for item in spec.get("images", []):
                    src, ratio = self._image_source(item["file_id"])
                    width = min(495, float(item.get("width", 495)))
                    if width <= 0:
                        raise FileSkillError("PDF图片宽度无效")
                    height = width / ratio
                    if height > 730:
                        width, height = 730 * ratio, 730
                    if y + height > 790:
                        new_page()
                    page.insert_image(fitz.Rect(50, y, 50 + width, y + height), filename=str(src))
                    y += height + 12
            if all(not p.get_text().strip() and not p.get_images() for p in pdf):
                raise FileSkillError("PDF内容不能为空")
            pdf.subset_fonts()
            pdf.save(path, garbage=4, deflate=True)
        finally:
            pdf.close()

    def edit_file(self, file_id, filename, operations):
        src = self._source(file_id)
        ext = src.suffix.lower()[1:]
        path, filename = self._target(filename, ext)
        if not isinstance(operations, list) or not 1 <= len(operations) <= 200:
            raise FileSkillError("须提供1至200项具体修改操作")
        for operation in operations:
            kind = operation.get("op") if isinstance(operation, dict) else None
            if kind not in _EDIT_FIELDS.get(ext, {}):
                raise FileSkillError("此格式不支持该修改操作")
            unsupported = set(operation) - {"op"} - set(_EDIT_FIELDS[ext][kind].split())
            if unsupported:
                raise FileSkillError("修改参数不受支持，不能忽略指定要求：" + ", ".join(sorted(unsupported)))
        try:
            if ext == "docx":
                doc = Document(src)
                for op in operations:
                    kind = op.get("op")
                    if kind == "replace_text":
                        old, new = _text(op["old"]), _text(op["new"])
                        count = sum(_replace_paragraph(p, old, new) for p in _docx_paragraphs(doc))
                        if not count:
                            raise FileSkillError("Word中未找到指定替换原文")
                    elif kind == "set_paragraph":
                        p = doc.paragraphs[_index(op["index"], len(doc.paragraphs), "段落")]
                        if p.text:
                            if not _replace_paragraph(p, p.text, _text(op["text"])):
                                raise FileSkillError("该段落含复杂链接或公式，不能安全整段替换；请指定可匹配原文")
                        else:
                            p.add_run(_text(op["text"]))
                    elif kind == "set_table_cell":
                        table = doc.tables[_index(op["table"], len(doc.tables), "表格")]
                        row = table.rows[_index(op["row"], len(table.rows), "行")]
                        cell = row.cells[_index(op["col"], len(row.cells), "列")]
                        if cell.tables or len(cell.paragraphs) > 1:
                            raise FileSkillError("复杂单元格请使用replace_text以保留内部结构")
                        p = cell.paragraphs[0]
                        if p.text:
                            if not _replace_paragraph(p, p.text, _text(op["text"])):
                                raise FileSkillError("复杂单元格无法安全替换，请指定可匹配原文")
                        else:
                            p.add_run(_text(op["text"]))
                    elif kind == "add_paragraph":
                        doc.add_paragraph(_text(op["text"]), style=op.get("style"))
                    elif kind == "add_table":
                        self._docx_table(doc, op)
                    elif kind == "add_image":
                        self._docx_image(doc, op)
                    elif kind in {"replace_image", "remove_image"}:
                        image = doc.inline_shapes[_index(op["index"], len(doc.inline_shapes), "图片")]
                        if kind == "remove_image":
                            image._inline.getparent().remove(image._inline)
                        else:
                            image_path, _ratio = self._image_source(op["file_id"])
                            relationship, _image = doc.part.get_or_add_image(str(image_path))
                            image._inline.graphic.graphicData.pic.blipFill.blip.embed = relationship
                    elif kind == "add_equation":
                        equation = OxmlElement("m:oMath")
                        run = OxmlElement("m:r")
                        text = OxmlElement("m:t")
                        text.text = _text(op["text"])
                        run.append(text)
                        equation.append(run)
                        doc.add_paragraph()._p.append(equation)
                    elif kind == "style_paragraph":
                        self._docx_style(doc.paragraphs[_index(op["index"], len(doc.paragraphs))], op)
                    else:
                        raise FileSkillError("Word修改操作不支持")
                doc.save(path)
            elif ext == "pptx":
                prs = Presentation(src)
                for op in operations:
                    kind = op.get("op")
                    if kind == "add_slide":
                        self._pptx_slide(prs, op)
                        continue
                    if kind == "replace_text":
                        old, new = _text(op["old"]), _text(op["new"])
                        count = 0
                        for slide in prs.slides:
                            for shape in slide.shapes:
                                frames = [shape.text_frame] if shape.has_text_frame else []
                                if shape.has_table:
                                    frames.extend(c.text_frame for row in shape.table.rows for c in row.cells)
                                for frame in frames:
                                    for p in frame.paragraphs:
                                        count += _replace_paragraph(p, old, new)
                        if not count:
                            raise FileSkillError("PPT中未找到指定原文")
                        continue
                    slide = prs.slides[_index(op["slide"], len(prs.slides), "幻灯片")]
                    if kind == "add_image":
                        image_path, ratio = self._image_source(op["file_id"])
                        x, y = float(op.get("x", .65)), float(op.get("y", 1.4))
                        width = float(op.get("width", 5))
                        height = float(op.get("height", width / ratio))
                        if min(x, y) < 0 or min(width, height) <= 0 or x + width > prs.slide_width / 914400 or y + height > prs.slide_height / 914400:
                            raise FileSkillError("新增图片超出幻灯片页面")
                        slide.shapes.add_picture(str(image_path), Inches(x), Inches(y), width=Inches(width), height=Inches(height))
                        continue
                    shape = slide.shapes[_index(op["shape"], len(slide.shapes), "图形")]
                    if kind == "set_shape_text":
                        if not shape.has_text_frame:
                            raise FileSkillError("图形没有文字")
                        if len(shape.text_frame.paragraphs) > 1:
                            raise FileSkillError("多段文字请使用replace_text保留其格式")
                        p = shape.text_frame.paragraphs[0]
                        if p.text:
                            if not _replace_paragraph(p, p.text, _text(op["text"])):
                                raise FileSkillError("PPT文字含复杂换行或域，无法安全整段替换，请指定可匹配原文")
                        else:
                            p.text = _text(op["text"])
                        self._pptx_check_text(shape)
                    elif kind == "set_table_cell":
                        if not shape.has_table:
                            raise FileSkillError("目标不是表格")
                        table = shape.table
                        cell = table.cell(_index(op["row"], len(table.rows)), _index(op["col"], len(table.columns)))
                        cell.text = _text(op["text"])
                    elif kind == "style_shape":
                        self._pptx_text_style(shape, op)
                        self._pptx_check_text(shape)
                    elif kind == "replace_image":
                        if not hasattr(shape, "_pic"):
                            raise FileSkillError("目标不是图片")
                        image_path, _ratio = self._image_source(op["file_id"])
                        _image_part, relationship = slide.part.get_or_add_image_part(str(image_path))
                        shape._pic.blipFill.blip.rEmbed = relationship
                    else:
                        raise FileSkillError("PPT修改操作不支持")
                prs.save(path)
            elif ext == "xlsx":
                wb = load_workbook(src, data_only=False, keep_links=False)
                try:
                    for op in operations:
                        self._xlsx_operation(wb, op)
                    wb.save(path)
                finally:
                    wb.close()
            elif ext == "pdf":
                with fitz.open(src) as pdf:
                    if pdf.needs_pass:
                        raise FileSkillError("PDF已加密")
                    for op in operations:
                        kind = op.get("op")
                        if kind == "delete_pages":
                            pages = op.get("pages", [])
                            for i in pages:
                                _index(i, len(pdf), "页")
                            if len(set(pages)) >= len(pdf):
                                raise FileSkillError("不能删除全部PDF页面")
                            pdf.delete_pages(sorted(set(pages)))
                            continue
                        page = pdf[_index(op["page"], len(pdf), "页")]
                        fontname, _font = _pdf_font(page, unique=True)
                        if kind == "replace_text":
                            old, new = _text(op["old"]), _text(op["new"])
                            if not old:
                                raise FileSkillError("替换原文不能为空")
                            rects = page.search_for(old)
                            if not rects:
                                raise FileSkillError("PDF中找不到指定原文；扫描PDF需OCR后重建")
                            size = float(op.get("font_size", max(6, min(12, rects[0].height * .65))))
                            for rect in rects:
                                page.add_redact_annot(rect, fill=(1, 1, 1))
                            page.apply_redactions(images=0, graphics=0)
                            # Redaction discards unused resources; bind the new font afterwards.
                            fontname, _font = _pdf_font(page, unique=True)
                            for rect in rects:
                                target = fitz.Rect(rect.x0, rect.y0 - 2, rect.x1, rect.y1 + 3)
                                if page.insert_textbox(target, new, fontname=fontname, fontsize=size) < 0:
                                    raise FileSkillError("替换文字放不下原PDF区域，请缩短文字或指定较小字号")
                        elif kind == "add_text":
                            size = float(op.get("font_size", 11))
                            x, y = float(op.get("x", 50)), float(op.get("y", 50))
                            if not 6 <= size <= 72 or not page.rect.contains(fitz.Point(x, y)):
                                raise FileSkillError("PDF文字位置或字号无效")
                            target = fitz.Rect(x, y, float(op.get("right", page.rect.width - 35)), float(op.get("bottom", page.rect.height - 35)))
                            if not page.rect.contains(target) or page.insert_textbox(target, _text(op["text"]), fontname=fontname, fontsize=size) < 0:
                                raise FileSkillError("新增PDF文字超出指定区域")
                        elif kind == "add_image":
                            image = self._source(op["file_id"])
                            rect = fitz.Rect(op["rect"])
                            if not page.rect.contains(rect) or rect.is_empty:
                                raise FileSkillError("图片区域超出PDF页面")
                            page.insert_image(rect, filename=str(image))
                        else:
                            raise FileSkillError("PDF修改操作不支持")
                    pdf.subset_fonts()
                    pdf.save(path, garbage=4, deflate=True)
            else:
                raise FileSkillError("此文件格式不支持直接编辑")
            return self._register(path, filename, f"已完成{len(operations)}项修改并重新打开检查")
        except Exception:
            path.unlink(missing_ok=True)
            raise

    def _xlsx_operation(self, wb, op):
        name = op.get("sheet")
        if name not in wb.sheetnames:
            raise FileSkillError("工作表不存在，请先读取文件")
        ws = wb[name]
        kind = op.get("op")
        def cells(address):
            if not isinstance(address, str) or not re.fullmatch(r"[A-Za-z]{1,3}[1-9]\d{0,6}(?::[A-Za-z]{1,3}[1-9]\d{0,6})?", address):
                raise FileSkillError("单元格地址无效")
            from openpyxl.utils.cell import range_boundaries
            min_col, min_row, max_col, max_row = range_boundaries(address)
            if max_col > 16384 or max_row > 1048576 or (max_col - min_col + 1) * (max_row - min_row + 1) > 100000:
                raise FileSkillError("单元格范围过大")
            return ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)
        if kind == "set_cell":
            if ":" in op["cell"]:
                raise FileSkillError("set_cell只接受单个单元格")
            rows = cells(op["cell"])
            next(rows)[0].value = _value(op.get("value"))
        elif kind == "set_range":
            start = op.get("start", "A1")
            if ":" in start:
                raise FileSkillError("起始位置须为单个单元格")
            base = next(cells(start))[0]
            values = op.get("values", [])
            if not values or len(values) > 10000 or sum(len(r) for r in values) > 100000:
                raise FileSkillError("写入范围过大或为空")
            for ri, row in enumerate(values):
                for ci, value in enumerate(row):
                    if base.row + ri > 1048576 or base.column + ci > 16384:
                        raise FileSkillError("写入范围超出Excel上限")
                    ws.cell(base.row + ri, base.column + ci, _value(value))
        elif kind == "style_range":
            for row in cells(op["range"]):
                for cell in row:
                    font = copy.copy(cell.font)
                    if "bold" in op:
                        font.bold = bool(op["bold"])
                    if "font_name" in op:
                        font.name = _text(op["font_name"])
                    if "font_size" in op:
                        size = float(op["font_size"])
                        if not 6 <= size <= 96:
                            raise FileSkillError("字号超出范围")
                        font.sz = size
                    if "color" in op:
                        font.color = _color(op["color"])
                    cell.font = font
                    if "fill" in op:
                        cell.fill = PatternFill("solid", fgColor=_color(op["fill"]))
                    if "number_format" in op:
                        cell.number_format = _text(op["number_format"])
                    if "wrap_text" in op:
                        alignment = copy.copy(cell.alignment)
                        alignment.wrap_text = bool(op["wrap_text"])
                        cell.alignment = alignment
        elif kind == "merge_cells":
            list(cells(op["range"]))
            ws.merge_cells(op["range"])
        elif kind == "column_width":
            column = str(op["column"]).upper()
            width = float(op["width"])
            if not re.fullmatch(r"[A-Z]{1,3}", column) or not 1 <= width <= 100:
                raise FileSkillError("列宽无效")
            ws.column_dimensions[column].width = width
        elif kind in {"add_image", "replace_image", "remove_image"}:
            if kind in {"replace_image", "remove_image"}:
                image = ws._images[_index(op["index"], len(ws._images), "图片")]
                if kind == "remove_image":
                    ws._images.remove(image)
                    return
            src, _ratio = self._image_source(op["file_id"])
            from openpyxl.drawing.image import Image as SheetImage
            new_image = SheetImage(str(src))
            if kind == "replace_image":
                new_image.anchor = copy.deepcopy(image.anchor)
                new_image.width, new_image.height = image.width, image.height
                ws._images[op["index"]] = new_image
            else:
                next(cells(op.get("cell", "A1")))
                if "width_px" in op:
                    new_image.width = max(1, min(2000, float(op["width_px"])))
                if "height_px" in op:
                    new_image.height = max(1, min(2000, float(op["height_px"])))
                ws.add_image(new_image, op.get("cell", "A1"))
        else:
            raise FileSkillError("Excel修改操作不支持")

    def validate_file(self, file_id):
        path = self._source(file_id)
        info = self.inspect_file(file_id)
        ext = path.suffix.lower()[1:]
        if path.stat().st_size == 0:
            raise FileSkillError("输出文件为空")
        if ext == "xlsx":
            wb = load_workbook(path, data_only=False, read_only=True, keep_links=False)
            try:
                for ws in wb:
                    if ws.max_row * ws.max_column > 300000:
                        raise FileSkillError("工作表尺寸超过处理限制")
                    for row in ws.iter_rows():
                        for cell in row:
                            if cell.data_type == "f":
                                _value(cell.value)
            finally:
                wb.close()
        elif ext == "pdf":
            with fitz.open(path) as pdf:
                if not len(pdf) or len(pdf) > 500:
                    raise FileSkillError("PDF页数无效或超出限制")
                for page in pdf:
                    if page.rect.is_empty:
                        raise FileSkillError("PDF包含无效页面")
        elif ext == "pptx":
            prs = Presentation(path)
            if not len(prs.slides):
                raise FileSkillError("PPT没有幻灯片")
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width + 1000 or shape.top + shape.height > prs.slide_height + 1000:
                        raise FileSkillError("幻灯片包含超出页面的图形，请调整布局")
                    if shape.has_text_frame:
                        self._pptx_check_text(shape)
        return {"valid": True, "checked": ["reopen", "structure", "size", "external_resources"], **info}
